#!/usr/bin/env python3
"""渲染算子表：YAML -> xlsx。

三个 sheet 对应 schema 的 rows / params / groups，有 accounting 时再加一张带宽账。
两种布局（--layout）：

review（默认）  核查用。后 7 列是组号 / 归属 / 核查结论 / 实现分叉 / 两侧证据 / 说明，
               证据锚点单独成列，便于逐行对源码。
                 existing        灰  已有算子承接，不排期
                 generic         蓝  任何同类模型都要，不该占本模型优先级
                 model-specific  绿  只有本模型的特殊结构才需要，真正要做的活
delivery       交付给算子团队。对齐既有家规的 Kimi 21 列：next3 / 并行资源占用 /
               duration / 计算利用率 / mte / scalar / 融合算子分析。组号与归属写进
               融合算子分析开头（【F7·本模型专属】），证据与分叉追加在同一格末尾。
               配色按组：已有=暖灰、通用=绿、专属=蓝，同色系三档深浅循环，相邻组不同色；
               组的上下边界加粗。schema 没有的列（next3、性能五列）恒为空 —— 没有实测不填估算。

用法:
  python3 scripts/mk_ops_table.py data-private/ops/<id>/table.yaml -o out.xlsx
  python3 scripts/mk_ops_table.py data-private/ops/<id>/table.yaml --layout delivery
  python3 scripts/mk_ops_table.py tests/fixtures/ops-table.sample.yaml   # 验流程
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lib import yamlio  # noqa: E402

COLUMNS = [
    ("module", "模块", 18),
    ("op", "算子列表", 30),
    ("kernel", "npu kernel", 34),
    ("op_type", "算子类型", 20),
    ("input_shape", "inputshape", 34),
    ("output_shape", "output shape", 24),
    ("dim", "dim", 10),
    ("shape_range", "shape范围", 34),
    ("input_dtype", "input dtype", 20),
    ("output_dtype", "output dtype", 16),
    ("input_format", "input format", 14),
    ("output_format", "output format", 14),
    ("dtype_next1", "dtype next1", 18),
    ("dtype_next2", "dtype next2", 18),
    ("group", "组号", 8),
    ("ownership", "归属", 14),
    ("_verdict", "核查结论", 12),
    ("_divergence", "实现分叉", 40),
    ("_impl", "实现侧证据", 30),
    ("_ref_impl", "参考实现证据", 30),
    ("note", "说明", 80),
]

FILL = {
    "existing": "D9D9D9",        # 灰：不排期
    "generic": "DCE6F1",         # 蓝：通用
    "model-specific": "D7E4BD",  # 绿：真正要做的
}
OWNERSHIP_CN = {
    "existing": "已有·无需新增",
    "generic": "通用",
    "model-specific": "本模型专属",
}
STATUS_CN = {
    "have": "已有", "extend": "扩展已有", "new": "要新写", "code-only": "纯代码改动",
}

# ---- delivery 布局：Kimi 21 列
# 第二项是行字段名；以 "_" 开头的是渲染时合成的。schema 里没有的字段（next3、性能五列）
# 读出来恒为空，列照样留着，算子团队拿到的表与家规格式一致。
DELIVERY_COLUMNS = [
    ("模块", "module", 26), ("算子列表", "op", 26), ("npu kernel", "kernel", 46),
    ("算子类型", "op_type", 30), ("inputshape", "input_shape", 46), ("output shape", "output_shape", 30),
    ("dim", "dim", 18), ("shape范围", "shape_range", 46), ("input dtype", "input_dtype", 30),
    ("output dtype", "output_dtype", 22), ("input format", "input_format", 30),
    ("output format", "output_format", 24), ("dtype next1", "dtype_next1", 30),
    ("dtype next2", "dtype_next2", 30), ("next3", "dtype_next3", 30),
    ("并行资源占用情况", "parallel_resource", 16), ("duration", "duration", 10),
    ("计算利用率", "compute_util", 12), ("mte", "mte", 8), ("scalar", "scalar", 8),
    ("融合算子分析", "_analysis", 70),
]
DELIVERY_GROUP_COLUMNS = [
    ("组号", 8), ("色块", 8), ("归属", 15), ("类型", 20), ("融合算子（已有 / 建议）", 54),
    ("收益 / 价值", 8), ("行数", 8), ("包含算子（表中相邻）", 64), ("说明", 90),
]
# 同色系三档深浅循环。已有=暖灰：一个组一旦已有，通用还是专属就不影响排期了
SHADES = {
    "existing": ["F0EFEA", "DEDCD3", "C8C5B8"],
    "generic": ["E2EFDA", "C6E0B4", "A9D08E"],
    "model-specific": ["DDEBF7", "BDD7EE", "9DC3E6"],
}
LEGEND = {
    "existing": "已有算子承接 —— 仓内 / 外部包 / 开关即可，不排期",
    "generic": "任何同类 TP transformer 都要 —— NPU 侧可能已有现成算子或图编译器能覆盖，不该占本模型的优先级",
    "model-specific": "只有本模型的特殊结构才需要 —— 真正要为它做的活",
}
DELIVERY_STATUS_CN = {
    "have": "已有", "extend": "扩展已有算子", "new": "新算子", "code-only": "纯代码改动",
}
BENEFIT_CN = {"high": "高", "medium": "中", "low": "低"}
HEADER_FILL = "D9E1F2"
BASELINE_KEYS = [
    ("hardware", "目标硬件"), ("soc_version", "SOC 版本"), ("parallel", "并行配置"),
    ("exec_mode", "执行模式"), ("launch_args", "启动参数"), ("quantization", "量化"),
    ("block_size", "block size"), ("config_caveat", "配置来源"),
]


def anchor_str(a):
    if not a:
        return ""
    s = a.get("file", "")
    if a.get("repo"):
        s = f"{a['repo']}:{s}"
    anc = (a.get("anchor") or "").strip().splitlines()
    return f"{s}\n@ {anc[0][:60]}" if anc else s


def flatten(row):
    ev = row.get("evidence") or {}
    out = dict(row)
    out["_verdict"] = ev.get("verdict", "")
    out["_divergence"] = ev.get("divergence", "")
    out["_impl"] = anchor_str(ev.get("impl"))
    out["_ref_impl"] = anchor_str(ev.get("ref_impl"))
    out["ownership"] = OWNERSHIP_CN.get(row.get("ownership", ""), row.get("ownership", ""))
    # 核查留痕并进说明列，不覆盖原值
    rv = row.get("review") or []
    if rv:
        marks = "\n".join(f"【核查 {r['id']}·{r.get('resolution', 'open')}】{r['text']}" for r in rv)
        out["note"] = (out.get("note") or "") + ("\n" if out.get("note") else "") + marks
    return out


def build(doc, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="BFBFBF")
    wrap = Alignment(vertical="top", wrap_text=True)

    wb = Workbook()

    # ---- 算子表
    ws = wb.active
    ws.title = "算子表"
    for c, (_, label, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, label)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"

    rows = doc.get("rows") or []
    for r, raw in enumerate(rows, 2):
        row = flatten(raw)
        fill = FILL.get(raw.get("ownership"))
        for c, (key, _, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(r, c, row.get(key, ""))
            cell.alignment = wrap
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)

    # ---- 模型参数
    ws2 = wb.create_sheet("模型参数")
    for c, (label, width) in enumerate([("符号 / 条目", 34), ("含义与取值（★ = 容易静默出错）", 120)], 1):
        cell = ws2.cell(1, c, label)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
        ws2.column_dimensions[get_column_letter(c)].width = width
    ws2.freeze_panes = "A2"
    r = 2
    last_section = None
    for p in doc.get("params") or []:
        if p.get("section") and p["section"] != last_section:
            last_section = p["section"]
            cell = ws2.cell(r, 1, f"—— {last_section} ——")
            cell.font = head_font
            r += 1
        key = ("★ " if p.get("pitfall") else "") + p["key"]
        ws2.cell(r, 1, key).alignment = wrap
        ws2.cell(r, 2, str(p["value"])).alignment = wrap
        r += 1

    # ---- 融合分组
    ws3 = wb.create_sheet("融合分组")
    cols3 = [("组号", 8), ("归属", 16), ("状态", 14), ("融合算子（已有 / 建议）", 46),
             ("收益", 8), ("收益 / 价值说明", 60), ("行数", 7), ("包含算子", 50), ("说明", 90)]
    for c, (label, width) in enumerate(cols3, 1):
        cell = ws3.cell(1, c, label)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
        ws3.column_dimensions[get_column_letter(c)].width = width
    ws3.freeze_panes = "A2"
    for r, g in enumerate(doc.get("groups") or [], 2):
        members = g.get("members") or []
        vals = [g["id"], OWNERSHIP_CN.get(g["ownership"], g["ownership"]),
                STATUS_CN.get(g.get("status", ""), g.get("status", "")),
                g.get("fused_op", ""), g.get("benefit", ""), g.get("benefit_note", ""),
                len(members), " → ".join(members), g.get("note", "")]
        fill = FILL.get(g.get("ownership"))
        for c, v in enumerate(vals, 1):
            cell = ws3.cell(r, c, v)
            cell.alignment = wrap
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)

    # 图例：配色的语义比配色本身重要，写在表里省得每次解释
    r = (len(doc.get("groups") or [])) + 3
    ws3.cell(r, 1, "配色 = 要不要干活").font = head_font
    for i, (k, cn) in enumerate(OWNERSHIP_CN.items(), 1):
        n = sum(1 for g in (doc.get("groups") or []) if g.get("ownership") == k)
        rowsn = sum(1 for x in rows if x.get("ownership") == k)
        cell = ws3.cell(r + i, 1, cn)
        cell.fill = PatternFill("solid", fgColor=FILL[k])
        ws3.cell(r + i, 2, f"{n} 组 / {rowsn} 行")

    # ---- 带宽账
    acc = doc.get("accounting") or []
    if acc:
        ws4 = wb.create_sheet("带宽账")
        for c, (label, width) in enumerate(
                [("条目", 34), ("取值", 22), ("依据", 12), ("与官方口径差异", 60), ("可修之处", 50)], 1):
            cell = ws4.cell(1, c, label)
            cell.font, cell.fill, cell.alignment = head_font, head_fill, wrap
            ws4.column_dimensions[get_column_letter(c)].width = width
        for r, a in enumerate(acc, 2):
            for c, v in enumerate([a["item"], str(a["value"]), a["basis"],
                                   a.get("vs_official", ""), a.get("fixable", "")], 1):
                ws4.cell(r, c, v).alignment = wrap

    wb.save(path)


def group_of(doc):
    """算子名 -> 组。以 groups[].members 为准。"""
    return {m: g for g in doc.get("groups") or [] for m in g.get("members") or []}


def group_colors(doc):
    """组号 -> 色值。按组在表中首次出现的顺序，在本归属的色系里循环取档；
    与表中紧挨着的上一组撞色时顺延一档，所以相邻组必不同色。"""
    gof = group_of(doc)
    cursor = {k: 0 for k in SHADES}
    colors, prev = {}, None
    for r in doc.get("rows") or []:
        g = gof.get(r.get("op"))
        if g is None:
            prev = None
            continue
        if g["id"] not in colors:
            shades = SHADES.get(g["ownership"], ["FFFFFF"])
            rgb = shades[cursor.get(g["ownership"], 0) % len(shades)]
            if prev is not None and colors.get(prev) == rgb:
                cursor[g["ownership"]] += 1
                rgb = shades[cursor[g["ownership"]] % len(shades)]
            cursor[g["ownership"]] = cursor.get(g["ownership"], 0) + 1
            colors[g["id"]] = rgb
        prev = g["id"]
    return colors


def analysis_text(row, g):
    """融合算子分析列：【组号·归属】+ 说明 + 核查留痕 + 分叉 + 两侧证据（标明哪一侧）。"""
    parts = []
    head = f"【{g['id']}·{OWNERSHIP_CN.get(g['ownership'], g['ownership'])}】" if g else ""
    parts.append(head + (row.get("note") or ""))
    for rv in row.get("review") or []:
        parts.append(f"★★【核查 {rv['id']}·{rv.get('resolution', 'open')}】{rv['text']}")
    ev = row.get("evidence") or {}
    if ev.get("divergence"):
        parts.append(f"★【实现分叉】{ev['divergence']}")
    refs = []
    for key, side in (("impl", "实现侧"), ("ref_impl", "参考实现"), ("upstream", "上游")):
        a = ev.get(key)
        if a:
            refs.append(f"{side} " + anchor_str(a).replace("\n", " "))
    if refs:
        verdict = f"（{ev['verdict']}）" if ev.get("verdict") else ""
        parts.append(f"〔证据{verdict}〕" + "；".join(refs))
    return "\n".join(p for p in parts if p)


def build_delivery(doc, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    hfill = PatternFill("solid", fgColor=HEADER_FILL)
    wrap = Alignment(vertical="top", wrap_text=True)
    thin, med = Side(style="thin", color="A6A6A6"), Side(style="medium", color="595959")

    def frame(ws, ncols, first, last, tops=(), bots=()):
        """全表细框、外框加粗；表头与 tops/bots 指定的行加粗上/下边。"""
        for r in range(first, last + 1):
            for c in range(1, ncols + 1):
                ws.cell(r, c).border = Border(
                    left=med if c == 1 else thin, right=med if c == ncols else thin,
                    top=med if (r == first or r in tops) else thin,
                    bottom=med if (r == first or r == last or r in bots) else thin)

    def header(ws, row, cols):
        for c, (label, width) in enumerate(cols, 1):
            cell = ws.cell(row, c, label)
            cell.font, cell.fill = bold, hfill
            ws.column_dimensions[get_column_letter(c)].width = width

    rows = doc.get("rows") or []
    groups = doc.get("groups") or []
    gof = group_of(doc)
    colors = group_colors(doc)

    wb = Workbook()

    # ---- 算子表
    ws = wb.active
    ws.title = "算子表"
    header(ws, 1, [(label, width) for label, _, width in DELIVERY_COLUMNS])
    ws.freeze_panes = "A2"
    tops, bots = set(), set()
    for r, row in enumerate(rows, 2):
        g = gof.get(row.get("op"))
        fill = PatternFill("solid", fgColor=colors[g["id"]]) if g else None
        for c, (_, key, _) in enumerate(DELIVERY_COLUMNS, 1):
            v = analysis_text(row, g) if key == "_analysis" else row.get(key, "")
            cell = ws.cell(r, c, v if v != "" else None)
            cell.alignment = wrap
            if fill:
                cell.fill = fill
        prev_g = gof.get(rows[r - 3].get("op")) if r > 2 else None
        next_g = gof.get(rows[r - 1].get("op")) if r - 1 < len(rows) else None
        if g and g is not prev_g:
            tops.add(r)
        if g and g is not next_g:
            bots.add(r)
    frame(ws, len(DELIVERY_COLUMNS), 1, len(rows) + 1, tops, bots)

    # ---- 模型参数：基线块在前，open_items 收在末尾
    ws2 = wb.create_sheet("模型参数")
    header(ws2, 1, [("符号 / 条目", 28), ("含义与取值（★ = 容易静默出错，重点看）", 110)])
    ws2.freeze_panes = "A2"
    base = doc.get("baseline") or {}
    entries = [{"section": "基线配置", "key": label, "value": str(base[k])}
               for k, label in BASELINE_KEYS if base.get(k) not in (None, "")]
    entries += doc.get("params") or []
    entries += [{"section": "未决与待实测", "key": str(i), "value": t}
                for i, t in enumerate(doc.get("open_items") or [], 1)]
    r, last = 2, None
    for p in entries:
        if p.get("section") and p["section"] != last:
            last = p["section"]
            ws2.cell(r, 1, f"—— {last} ——")
            r += 1
        key = p["key"]
        if p.get("pitfall") and not key.startswith("★"):
            key = "★" + key
        ws2.cell(r, 1, key).alignment = wrap
        ws2.cell(r, 2, str(p["value"])).alignment = wrap
        r += 1
    frame(ws2, 2, 1, r - 1)

    # ---- 融合分组：图例在上，组表从第 8 行起
    ws3 = wb.create_sheet("融合分组")
    ws3.cell(1, 1, "配色 = 要不要干活").font = bold
    for i, k in enumerate(SHADES, 2):
        ng = sum(1 for g in groups if g["ownership"] == k)
        nr = sum(len(g.get("members") or []) for g in groups if g["ownership"] == k)
        ws3.cell(i, 2).fill = PatternFill("solid", fgColor=SHADES[k][0])
        ws3.cell(i, 3, OWNERSHIP_CN[k]).font = bold
        ws3.cell(i, 4, f"{ng} 组 / {nr} 行")
        ws3.cell(i, 5, LEGEND[k])
    standalone = len(rows) - len(gof)
    ws3.cell(6, 1, "★两列怎么读").font = bold
    ws3.cell(6, 4, "「融合算子」列")
    ws3.cell(6, 5, "灰色组 = 已有 →，即现成算子叫什么、已经融好了、不用排期；绿/蓝组 = 建议 →，"
                   "即建议新建或扩展成什么。收益列对灰色组读作「该融合的价值」，对绿/蓝组读作「做了的预期收益」。"
                   f"不分组的 {standalone} 行是单个算子即可、无相邻可融或控制流。")
    header(ws3, 8, DELIVERY_GROUP_COLUMNS)
    ws3.freeze_panes = "A9"
    for r, g in enumerate(groups, 9):
        members = g.get("members") or []
        note = g.get("note", "")
        if g.get("benefit_note"):
            note = f"【收益】{g['benefit_note']}" + (f"\n{note}" if note else "")
        vals = [g["id"], None, OWNERSHIP_CN.get(g["ownership"], g["ownership"]),
                DELIVERY_STATUS_CN.get(g.get("status", ""), g.get("status", "")),
                g.get("fused_op", ""), BENEFIT_CN.get(g.get("benefit", ""), g.get("benefit", "")),
                len(members), " → ".join(members), note]
        for c, v in enumerate(vals, 1):
            ws3.cell(r, c, v).alignment = wrap
        if g["id"] in colors:
            ws3.cell(r, 2).fill = PatternFill("solid", fgColor=colors[g["id"]])
    frame(ws3, len(DELIVERY_GROUP_COLUMNS), 8, 8 + len(groups))

    # ---- 带宽账（与 review 布局相同）
    acc = doc.get("accounting") or []
    if acc:
        ws4 = wb.create_sheet("带宽账")
        header(ws4, 1, [("条目", 34), ("取值", 22), ("依据", 12), ("与官方口径差异", 60), ("可修之处", 50)])
        for r, a in enumerate(acc, 2):
            for c, v in enumerate([a["item"], str(a["value"]), a["basis"],
                                   a.get("vs_official", ""), a.get("fixable", "")], 1):
                ws4.cell(r, c, v).alignment = wrap
        frame(ws4, 5, 1, len(acc) + 1)

    wb.save(path)


def check(doc):
    """渲染前的自查：分组引用的算子必须存在且在表中相邻，组内归属一致，行上的 group 与分组表一致。

    不写 group 的行是有意不分组（单个算子即可、无相邻可融）；写了 group 却不在那个组的
    members 里，才是漏登。"""
    errs = []
    rows = doc.get("rows") or []
    ops = {r.get("op") for r in rows}
    dup = len(rows) - len(ops)
    if dup > 0:
        errs.append(f"有 {dup} 个重复的算子名，分组引用会指向错的行")
    pos = {r.get("op"): i for i, r in enumerate(rows)}
    for g in doc.get("groups") or []:
        members = g.get("members") or []
        for m in members:
            if m not in ops:
                errs.append(f"分组 {g['id']} 引用了不存在的算子 '{m}'")
        idx = [pos[m] for m in members if m in pos]
        if idx and idx != list(range(idx[0], idx[0] + len(idx))):
            errs.append(f"分组 {g['id']} 的成员在表中不相邻（行 {[i + 2 for i in idx]}）——"
                        "融合算子只能覆盖执行序上连续的一段")
        owners = {r.get("ownership") for r in rows if r.get("op") in members}
        if len(owners) > 1:
            errs.append(f"分组 {g['id']} 内归属不一致：{owners}，配色会自相矛盾")
    gof = group_of(doc)
    for r in rows:
        g = gof.get(r.get("op"))
        if r.get("group") and (g is None or g["id"] != r["group"]):
            errs.append(f"算子 '{r.get('op')}' 标了 {r['group']}，但不在该组 members 里"
                        + (f"（实际在 {g['id']}）" if g else ""))
        elif g and r.get("group") is None:
            errs.append(f"算子 '{r.get('op')}' 在 {g['id']} 的 members 里，但行上没写 group")
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("yaml_path")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--force", action="store_true", help="自查有问题也照样渲染")
    ap.add_argument("--layout", choices=["review", "delivery"], default="review",
                    help="review=核查用（证据单独成列）；delivery=交付用的 Kimi 21 列")
    args = ap.parse_args()

    doc = yamlio.load(args.yaml_path)
    errs = check(doc)
    for e in errs:
        print(f"  ERROR {e}", file=sys.stderr)
    if errs and not args.force:
        sys.exit("! 自查未通过，修好再渲染（--force 可跳过）")

    out = args.out or os.path.join(
        os.path.dirname(args.yaml_path) or ".",
        f"{doc.get('model', 'ops')}_decode_算子表.xlsx")
    (build_delivery if args.layout == "delivery" else build)(doc, out)
    rows = doc.get("rows") or []
    by = {}
    for r in rows:
        by[r.get("ownership", "?")] = by.get(r.get("ownership", "?"), 0) + 1
    print(f"-> {out}")
    print(f"   {len(rows)} 行, {len(doc.get('groups') or [])} 组, "
          f"{len(doc.get('params') or [])} 条参数")
    print("   归属分布: " + ", ".join(f"{OWNERSHIP_CN.get(k, k)} {v}" for k, v in sorted(by.items())))
    base = doc.get("baseline") or {}
    if base.get("config_origin") == "assembled":
        print(f"   ! 基线是拼装配置: {base.get('config_caveat', '未说明风险')}")


if __name__ == "__main__":
    main()
